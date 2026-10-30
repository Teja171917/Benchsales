"""Core logic tests (stdlib unittest; no network, no web framework needed).

Run from the project root:   python -m unittest discover -s tests -v
"""
import os
import sqlite3
import sys
import json
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app import autolearn, contact, db, emptype, market, matcher, office, portals, rules, skills, tailor, usa  # noqa: E402
from app.sources import urlimport  # noqa: E402


class EmpTypeTests(unittest.TestCase):
    def check(self, et, title, desc, expected):
        self.assertEqual(set(emptype.classify(et, title, desc)), set(expected),
                         f"{et!r} {title!r} {desc!r}")

    def test_aggregator_type_words(self):
        self.check("full_time", "Java Developer", "", ["fulltime"])
        self.check("CONTRACTOR", "QA Engineer", "", ["contract"])
        self.check("part_time", "Dev", "", ["parttime"])

    def test_c2c_w2_in_description(self):
        self.check("", "Java Dev", "Open to C2C or W2. Hybrid in Dallas, TX",
                   ["c2c", "w2", "contract"])
        self.check("", "Dev", "Corp-to-Corp, 1099 and W2 fine, 12 month contract",
                   ["c2c", "w2", "1099", "contract"])
        self.check("", "Sr Python Dev (C2C)", "", ["c2c", "contract"])

    def test_negations(self):
        self.check("", "Dev", "W2 only. No C2C. No third party.", ["w2"])
        self.check("", "Dev", "C2C not accepted. W2 candidates only", ["w2"])
        self.check("", "Dev", "C2C is not allowed", [])
        self.check("", "Dev", "No corp to corp please; W2 contract role",
                   ["w2", "contract"])

    def test_real_posting_no_third_party_list(self):
        # real Adzuna/Columbus posting that was wrongly tagged C2C + 1099
        desc = ("$55/hr W2 Contract. USC/GC/GC EAD Only- W2 Requirement\n\n"
                "Note: No third party/C2C or 1099\n\nJob Overview")
        self.check("contract", "Java Developer", desc, ["w2", "contract"])
        self.check("", "Dev", "Open to C2C or W2, no 1099", ["c2c", "w2", "contract"])
        self.check("", "Dev", "No sponsorship, C2C ok", ["c2c", "contract"])
        self.check("", "Dev", "W2 required", ["w2"])

    def test_contract_to_hire(self):
        self.check("", "Data Eng", "Contract to hire, W2", ["w2", "contract", "c2h"])

    def test_boilerplate_is_not_a_type(self):
        self.check("", "Data Eng", "Great benefits for full-time employees", [])
        self.check("", "Contracts Manager", "manage vendor agreements", [])

    def test_filter_semantics(self):
        self.assertTrue(emptype.matches_filter(["c2c"], ["c2c", "w2"]))
        self.assertFalse(emptype.matches_filter(["fulltime"], ["c2c", "w2"]))
        self.assertFalse(emptype.matches_filter([], ["c2c"]))
        self.assertTrue(emptype.matches_filter([], ["c2c"], include_unspecified=True))
        self.assertTrue(emptype.matches_filter(["parttime"], []))


class UsaTests(unittest.TestCase):
    def test_locations(self):
        yes = ["Houston, TX", "Remote", "Worldwide", "", "Austin, Texas",
               "Remote - USA", "Americas", "Dallas-Fort Worth, TX",
               "Remote (Canada or US)", "Vienna, VA", "Melbourne, FL",
               "Boston (remote)", "Europe, USA, Canada, APAC (remote)"]
        no = ["Berlin, Germany", "Europe", "London, UK", "Toronto, Ontario, Canada",
              "Remote, India", "Bengaluru, Karnataka, India",
              "Bishkek, Bishkek City, Kyrgyzstan (remote)",
              "Kuala Lumpur, Malaysia (remote)", "Seoul (remote)",
              "Budapest, (remote)", "Melbourne (remote)",
              "Uluberia-II, (remote)"]
        for loc in yes:
            self.assertTrue(usa.is_us_job({"location": loc}), loc)
        for loc in no:
            self.assertFalse(usa.is_us_job({"location": loc}), loc)


class PortalTests(unittest.TestCase):
    def test_links(self):
        links = {l["id"]: l["url"] for l in
                 portals.build_links("Java Developer", "Houston, TX",
                                     ["fulltime", "c2c", "w2"], 7)}
        for pid in ("dice", "indeed", "linkedin", "glassdoor", "monster", "ziprecruiter"):
            self.assertIn(pid, links)
            self.assertTrue(links[pid].startswith("https://"))
        self.assertIn("Java+Developer", links["dice"])
        self.assertIn("THIRD_PARTY", links["dice"])
        self.assertEqual(links["dice"].count("filters.employmentType="), 1)
        self.assertIn("f_JT=F%2CC", links["linkedin"])
        self.assertIn("C2C", links["indeed"])

    def test_blank_location_means_united_states(self):
        links = {l["id"]: l["url"] for l in portals.build_links("qa", "", ["c2c"], 7)}
        for pid in ("indeed", "linkedin", "ziprecruiter", "monster"):
            self.assertIn("United", links[pid], pid)
        self.assertIn("united-states-qa-jobs-SRCH_IL.0,13_IN1_KO14,16", links["glassdoor"])
        usa = {l["id"]: l["url"] for l in portals.build_links("QA Engineer", "USA", [], 7)}
        self.assertIn("united-states-qa-engineer-jobs-SRCH_IL.0,13_IN1_KO14,25", usa["glassdoor"])
        self.assertIn("l=United+States", usa["indeed"])
        self.assertNotIn("C2C", links["ziprecruiter"])
        self.assertNotIn("lensa", links)
        self.assertNotIn("careerbuilder", links)


class DbTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._old = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "t.db"
        db.init_db()

    def tearDown(self):
        db.DB_PATH = self._old
        self.tmp.cleanup()

    def job(self, n, **kw):
        j = {"source": "t", "source_id": f"id{n}", "title": f"Java Dev {n}",
             "company": f"Co{n}", "location": "Houston, TX", "remote_flag": 0,
             "url": "", "description": "Java Spring Boot", "posted_at": "",
             "salary": "", "employment_type": ""}
        j.update(kw)
        return j

    def test_tags_stored_and_filtered(self):
        db.insert_job(self.job(1, employment_type="full_time"))
        db.insert_job(self.job(2, description="C2C or W2 contract"))
        db.insert_job(self.job(3, description="W2 only. No C2C."))
        db.insert_job(self.job(4))  # nothing detectable
        titles = lambda **kw: sorted(j["title"] for j in db.list_jobs(**kw))
        self.assertEqual(titles(emp=["c2c"]), ["Java Dev 2"])
        self.assertEqual(titles(emp=["w2"]), ["Java Dev 2", "Java Dev 3"])
        self.assertEqual(titles(emp=["fulltime", "c2c"]), ["Java Dev 1", "Java Dev 2"])
        self.assertEqual(titles(emp=["c2c"], include_unspecified=True),
                         ["Java Dev 2", "Java Dev 4"])
        self.assertEqual(len(db.list_jobs()), 4)
        self.assertIn("c2c", db.list_jobs(emp=["c2c"])[0]["emp_tags"])

    def test_stale_tags_are_retagged_on_start(self):
        db.insert_job(self.job(1, description="W2 Requirement. Note: No third party/C2C or 1099"))
        with db.get_conn() as c:   # simulate tags saved by the old buggy rules
            c.execute("UPDATE jobs SET emp_tags=',c2c,w2,1099,'")
            c.execute("DELETE FROM settings WHERE key='emp_tags_version'")
            c.commit()
        db.init_db()
        self.assertEqual(db.list_jobs()[0]["emp_tags"], ["w2"])

    def test_dedupe_uses_location(self):
        self.assertTrue(db.insert_job(self.job(1, title="Java Dev", company="Acme")))
        # same title+company+location from another source -> duplicate
        self.assertIsNone(db.insert_job(self.job(2, title="java dev", company="ACME")))
        # same title+company, different city -> a separate opening
        self.assertTrue(db.insert_job(self.job(3, title="Java Dev", company="Acme",
                                               location="Dallas, TX")))

    def test_migration_from_old_schema(self):
        p = Path(self.tmp.name) / "old.db"
        con = sqlite3.connect(p)
        con.executescript("""
        CREATE TABLE consultants (id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
          email TEXT DEFAULT '', phone TEXT DEFAULT '', location TEXT DEFAULT '',
          visa_status TEXT DEFAULT '', linkedin_url TEXT DEFAULT '', notes TEXT DEFAULT '',
          created_at TEXT NOT NULL);
        CREATE TABLE jobs (id INTEGER PRIMARY KEY AUTOINCREMENT, source TEXT NOT NULL,
          source_id TEXT NOT NULL, title TEXT DEFAULT '', company TEXT DEFAULT '',
          location TEXT DEFAULT '', remote_flag INTEGER DEFAULT 0, url TEXT DEFAULT '',
          description TEXT DEFAULT '', posted_at TEXT DEFAULT '', salary TEXT DEFAULT '',
          employment_type TEXT DEFAULT '', fetched_at TEXT NOT NULL, UNIQUE(source, source_id));
        INSERT INTO jobs(source, source_id, title, description, fetched_at)
          VALUES ('x','1','Old job','Corp to corp welcome','2026-01-01');
        """)
        con.commit()
        con.close()
        db.DB_PATH = p
        db.init_db()
        db.init_db()  # idempotent
        j = db.list_jobs(emp=["c2c"])
        self.assertEqual([x["title"] for x in j], ["Old job"])
        c = db.create_consultant({"name": "A", "emp_pref": "w2, c2c ,bogus"})
        self.assertEqual(c["emp_pref"], "w2,c2c")

    def test_matcher_respects_consultant_preference(self):
        w2_only = {"emp_pref": "w2"}
        self.assertTrue(matcher.emp_compatible(w2_only, ["w2", "c2c"]))
        self.assertFalse(matcher.emp_compatible(w2_only, ["c2c"]))
        self.assertTrue(matcher.emp_compatible(w2_only, []))      # unknown passes
        self.assertTrue(matcher.emp_compatible({"emp_pref": ""}, ["c2c"]))  # no pref

    def test_run_all_end_to_end(self):
        c1 = db.create_consultant({"name": "W2 Person", "emp_pref": "w2"})
        c2 = db.create_consultant({"name": "Any Person"})
        txt = "Senior Java Developer\nJava Spring Boot microservices REST API AWS Docker PostgreSQL"
        for c in (c1, c2):
            db.upsert_resume(c["id"], "r.txt", txt, matcher.extract_skills(txt))
        desc = ("Requirements:\n- Java, Spring Boot, microservices, REST API\n"
                "- AWS, Docker, PostgreSQL\n")
        db.insert_job(self.job(1, title="Senior Java Developer",
                               description=desc + "C2C only"))
        db.insert_job(self.job(2, title="Senior Java Developer", company="Other",
                               description=desc + "W2 only"))
        n = matcher.run_all(threshold=10)
        by_c = {}
        for m in db.list_matches():
            by_c.setdefault(m["consultant_name"], []).append(m["emp_tags"])
        self.assertEqual(len(by_c["Any Person"]), 2)
        self.assertEqual(by_c["W2 Person"], [["w2"]])  # C2C-only job skipped
        self.assertEqual(n, 3)
        self.assertEqual(matcher.run_all(threshold=10), 0)  # idempotent


    def test_run_all_for_one_consultant(self):
        c1 = db.create_consultant({"name": "A"})
        c2 = db.create_consultant({"name": "B"})
        txt = "Senior Java Developer\nJava Spring Boot microservices REST API AWS Docker"
        for c in (c1, c2):
            db.upsert_resume(c["id"], "r.txt", txt, matcher.extract_skills(txt))
        db.insert_job(self.job(1, description="Requirements:\n- Java, Spring Boot, "
                               "microservices, REST API, AWS, Docker\n"))
        self.assertEqual(matcher.run_all(threshold=10, consultant_id=c1["id"]), 1)
        self.assertEqual([m["consultant_name"] for m in db.list_matches()], ["A"])
        self.assertEqual(matcher.run_all(threshold=10, consultant_id=c1["id"]), 0)


class LocationScoreTests(unittest.TestCase):
    def sim(self, a, b):
        return matcher._jaccard(matcher._loc_tokens(a), matcher._loc_tokens(b))

    def test_state_abbreviation_matches_state_name(self):
        self.assertGreater(self.sim("California, USA", "Irvine, CA"), 0)
        self.assertEqual(self.sim("Irvine, CA", "Irvine, CA"), 1.0)
        self.assertEqual(self.sim("Houston, TX", "Irvine, CA"), 0.0)
        self.assertGreater(self.sim("New York, NY", "New York City"), 0.5)


class ResumeFileTests(unittest.TestCase):
    def make_docx(self):
        import io
        from docx import Document
        d = Document()
        d.add_paragraph("JANE DOE")
        d.add_paragraph("TECHNICAL SKILLS")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Languages:"
        t.cell(0, 1).text = "Java, Python"
        t.cell(1, 0).text = "Tools:"
        t.cell(1, 1).text = "Jenkins, Docker"
        d.add_paragraph("EXPERIENCE")
        d.add_paragraph("Built things at Acme")
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    def test_table_stays_where_it_is(self):
        from app import resume
        text = resume.parse("jane.docx", self.make_docx())
        lines = text.splitlines()
        self.assertEqual(lines[:2], ["JANE DOE", "TECHNICAL SKILLS"])
        self.assertEqual(lines[2], "Languages: Java, Python")
        self.assertEqual(lines[3], "Tools: Jenkins, Docker")
        self.assertEqual(lines[4], "EXPERIENCE")

    def test_txt_and_unsupported(self):
        from app import resume
        self.assertEqual(resume.parse("a.txt", b"hello"), "hello")
        with self.assertRaises(ValueError):
            resume.parse("a.exe", b"MZ")
        with self.assertRaises(Exception):
            resume.parse("bad.docx", b"not really a docx")


class TailorTests(unittest.TestCase):
    RESUME = ("JANE DOE\nQA Engineer\nSUMMARY\nQA person.\n\nTECHNICAL SKILLS\n"
              "Languages: Java, Python, SQL\n"
              "Tools: Jenkins, Docker, RestAssured, Selenium\n\n"
              "EXPERIENCE\nBuilt things.")

    def test_categorised_skills_keep_their_lines(self):
        from app import tailor
        skills = matcher.extract_skills(self.RESUME)
        out = tailor.keyword_tailor(self.RESUME, skills,
                                    "Requirements: RestAssured, Selenium, Docker")
        lines = out.splitlines()
        self.assertIn("Languages: Java, Python, SQL", lines)
        # JD skills move to the front of their own line, nothing is dropped
        self.assertIn("Tools: Docker, RestAssured, Selenium, Jenkins", lines)
        self.assertFalse(any(ln.startswith("Skills:") for ln in lines))

    def test_plain_skills_list_is_still_reordered(self):
        from app import tailor
        txt = "JANE DOE\nSKILLS\nJava, Python, Docker\n\nEXPERIENCE\nx"
        out = tailor.keyword_tailor(txt, ["java", "python", "docker"],
                                    "Requirements: Docker")
        self.assertIn("Skills: docker, java, python", out)


class SkillsTests(unittest.TestCase):
    def test_qa_and_ai_skills_found(self):
        from app import skills
        got = set(skills.extract_skills(
            "Frameworks: Rest Assured, Selenium WebDriver, SoapUI, Allure "
            "Reporting. AI: Claude API integration, Generative AI QA, "
            "ISTQB CT-AI."))
        for s in ("restassured", "selenium webdriver", "soapui", "allure",
                  "claude api", "generative ai", "istqb"):
            self.assertIn(s, got)

    def test_teams_word_is_not_microsoft_teams(self):
        from app import skills
        self.assertNotIn("microsoft teams", skills.extract_skills(
            "Led cross-functional teams of 14 engineers"))
        self.assertIn("microsoft teams", skills.extract_skills(
            "Daily standups on MS Teams"))

    def test_display_names(self):
        from app import skills
        self.assertEqual(skills.display_name("playwright"), "Playwright")
        self.assertEqual(skills.display_name("ci/cd"), "CI/CD")
        self.assertEqual(skills.display_name("typescript"), "TypeScript")
        self.assertEqual(skills.display_name("selenium webdriver"), "Selenium WebDriver")
        self.assertEqual(skills.display_name("test automation"), "Test Automation")

    def test_sentence_final_skill_still_matches(self):
        # regression: the "." guard used to swallow skills at sentence end
        from app import skills
        self.assertIn("playwright",
                      skills.extract_skills("Senior QA with Playwright."))
        self.assertIn("selenium",
                      skills.extract_skills("Hands-on Selenium."))
        # the guard's real job: no matching inside longer tokens
        self.assertNotIn("java", skills.extract_skills("JavaScript developer"))


class AutoTests(unittest.TestCase):
    """The 'works without humans' parts: learning, queries, quota, rescoring."""
    RESUME = ("Sharan Murali\nQA Automation Engineer | Dallas, TX\n\n"
              "SUMMARY\nTester.\n\nTECHNICAL SKILLS\n"
              "Languages: Java, Python, SQL\n"
              "Test Mgmt: Zephyr Scale, Xray (Jira), TestRail\n"
              "Tools: Selenium, Postman, Jenkins, and, etc\n\n"
              "EXPERIENCE\nWorked with Cobol Mainframe daily.\n")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._old = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "t.db"
        db.init_db()
        skills.set_learned([])

    def tearDown(self):
        skills.set_learned([])
        matcher._jd_skills.cache_clear()
        db.DB_PATH = self._old
        self.tmp.cleanup()

    def test_candidates_come_only_from_skills_section(self):
        got = autolearn.candidate_skills(self.RESUME)
        self.assertIn("zephyr scale", got)
        self.assertIn("xray", got)
        self.assertIn("testrail", got)
        for bad in ("and", "etc", "cobol mainframe", "java", "selenium"):
            self.assertNotIn(bad, got)

    def test_learning_changes_extraction_and_is_removable(self):
        c = db.create_consultant({"name": "S", "location": "Dallas, TX"})
        db.upsert_resume(c["id"], "r.txt", self.RESUME, skills.extract_skills(self.RESUME))
        self.assertNotIn("zephyr scale", db.get_resume(c["id"])["skills_json"])
        out = autolearn.maintain()
        self.assertIn("zephyr scale", out["learned"])
        self.assertEqual(out["resumes_refreshed"], 1)
        self.assertIn("zephyr scale", db.get_resume(c["id"])["skills_json"])
        self.assertIn("zephyr scale", skills.extract_skills("need Zephyr Scale"))
        self.assertEqual(autolearn.maintain()["learned"], [])  # nothing new 2nd time
        self.assertTrue(db.delete_learned_skill("zephyr scale"))
        autolearn.block_skill("zephyr scale")
        autolearn.maintain()   # must not learn it again from the same resume
        self.assertNotIn("zephyr scale", skills.extract_skills("need Zephyr Scale"))

    def test_rescoring_improves_scores_and_never_deletes(self):
        c = db.create_consultant({"name": "S"})
        db.upsert_resume(c["id"], "r.txt", self.RESUME, skills.extract_skills(self.RESUME))
        desc = ("Requirements:\n- Java, Selenium, Zephyr Scale, TestRail, Xray\n")
        jid = db.insert_job({"source": "t", "source_id": "1", "title": "QA Automation Engineer",
                             "company": "Co", "location": "Dallas, TX", "remote_flag": 0,
                             "url": "", "description": desc, "posted_at": "",
                             "salary": "", "employment_type": ""})
        self.assertEqual(matcher.run_all(threshold=0), 1)
        before = db.list_matches()[0]["score"]
        autolearn.maintain()  # learns the three tools -> job needs them, resume has them
        after = db.list_matches()
        self.assertEqual(len(after), 1)
        self.assertGreater(after[0]["score"], before)

    def test_queries_built_from_resumes(self):
        c1 = {"raw_text": self.RESUME, "skills": ["java"], "location": "Dallas, TX"}
        c2 = {"raw_text": "Ravi K\nJava Full Stack Developer\nJava Spring", "skills": ["java"],
              "location": ""}
        qs = autolearn.build_queries([c1, c2], [{"title": "Manual Role", "location": "Austin"}], 8)
        titles = [(q["title"], q["location"]) for q in qs]
        self.assertEqual(titles[0], ("QA Automation Engineer", "Dallas, TX"))
        self.assertIn(("Java Full Stack Developer", "Remote"), titles)
        self.assertIn(("Manual Role", "Austin"), titles)
        self.assertEqual(len(titles), len(set(titles)))
        self.assertEqual(len(autolearn.build_queries([c1, c2], [], 2)), 2)
        # nobody uploaded anything yet -> fall back to the typed queries
        self.assertEqual(autolearn.queries_for_run({"auto_queries": "1"}),
                         db.DEFAULT_SEARCH_QUERIES)

    def test_interval_respects_adzuna_daily_budget(self):
        base = {"collect_interval_minutes": "15", "enabled_sources": '["adzuna"]',
                "auto_queries": "0", "adzuna_daily_budget": "80"}
        # no keys -> no slowdown
        self.assertEqual(autolearn.effective_interval_minutes(base), 15)
        keyed = {**base, "adzuna_app_id": "a", "adzuna_app_key": "b"}
        n = len(autolearn.queries_for_run(keyed))        # 4 default queries -> 8 calls/run
        runs = 80 // (n * 2)
        self.assertEqual(autolearn.effective_interval_minutes(keyed), -(-1440 // runs))
        self.assertLessEqual(runs * n * 2, 80)
        self.assertEqual(autolearn.effective_interval_minutes({**keyed, "collect_interval_minutes": "0"}), 0)

    def test_collector_stops_at_daily_budget(self):
        import collector
        db.set_setting("adzuna_daily_budget", "6")
        qs = [{"title": f"T{i}", "location": ""} for i in range(5)]
        s = {"budget_skipped": {}}
        self.assertEqual(len(collector._within_budget(qs, s)), 3)   # 3 x 2 calls = 6
        self.assertEqual(s["budget_skipped"]["adzuna"], 2)
        db.usage_add("adzuna", 6)
        self.assertEqual(collector._within_budget(qs, {"budget_skipped": {}}), [])


class SsrfTests(unittest.TestCase):
    def test_private_hosts_blocked(self):
        for u in ["http://127.0.0.1:8741/api/settings",
                  "http://169.254.169.254/latest/meta-data/",
                  "http://localhost/", "file:///etc/passwd",
                  "http://10.0.0.5/", "http://[::1]/"]:
            with self.assertRaises(ValueError, msg=u):
                urlimport._assert_public(u)


class MarketTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._old = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "t.db"
        db.init_db()

    def tearDown(self):
        db.DB_PATH = self._old
        self.tmp.cleanup()

    def job(self, n, desc):
        return {"source": "t", "source_id": f"m{n}", "title": f"Dev {n}",
                "company": f"Co{n}", "location": "", "remote_flag": 0,
                "url": "", "description": desc, "posted_at": "",
                "salary": "", "employment_type": ""}

    def test_demand_ranks_by_posting_count(self):
        db.insert_job(self.job(1, "Need Playwright and TypeScript for UI automation"))
        db.insert_job(self.job(2, "Playwright expert wanted, CI/CD experience"))
        db.insert_job(self.job(3, "Java backend developer, Spring Boot"))
        top = market.demand(days=30, limit=10)
        by_skill = {d["skill"]: d["jobs"] for d in top}
        self.assertEqual(by_skill.get("playwright"), 2)
        self.assertEqual(by_skill.get("typescript"), 1)
        self.assertEqual(top[0]["skill"], "playwright")

    def test_demand_ignores_empty_descriptions(self):
        db.insert_job(self.job(1, ""))
        self.assertEqual(market.demand(days=30), [])


class AddSkillsTests(unittest.TestCase):
    def test_categorised_line_gets_new_skills(self):
        text = ("Jane Doe\n\nTechnical Skills\n"
                "Languages: Java, Python\n"
                "Tools: Jenkins, Docker\n"
                "\nExperience\n- Built stuff.")
        out = tailor.add_skills_to_text(text, ["Playwright", "TypeScript"])
        self.assertIn("Tools: Jenkins, Docker, Playwright, TypeScript", out)
        self.assertIn("Languages: Java, Python", out)

    def test_plain_block_gets_additional_line(self):
        text = "Jane Doe\n\nSkills\nJava, Python\n\nExperience\n- Built stuff."
        out = tailor.add_skills_to_text(text, ["Playwright"])
        self.assertIn("Additional skills: Playwright", out)

    def test_duplicates_skipped(self):
        text = "Jane Doe\n\nTechnical Skills\nLanguages: Java, Python\n\nExperience\n- Built stuff."
        out = tailor.add_skills_to_text(text, ["Java", "Playwright", "playwright"])
        self.assertEqual(out.count("Playwright"), 1)
        self.assertNotIn("playwright, Playwright", out)
        # Java was already present -> not added again
        self.assertEqual(out.count("Java"), 1)

    def test_no_skills_no_change(self):
        text = "Jane Doe\n\nSkills: Java"
        self.assertEqual(tailor.add_skills_to_text(text, []), text)

    def test_confirmed_skills_stored_and_unflagged(self):
        tmp = tempfile.TemporaryDirectory()
        old = db.DB_PATH
        db.DB_PATH = Path(tmp.name) / "t.db"
        try:
            db.init_db()
            c = db.create_consultant({"name": "T"})
            jid = db.insert_job({"source": "t", "source_id": "x1", "title": "Dev",
                                 "company": "Co", "location": "", "remote_flag": 0,
                                 "url": "", "description": "Playwright",
                                 "posted_at": "", "salary": "", "employment_type": ""})
            mid = db.insert_match(c["id"], jid, 65.0, {}, ["playwright"])
            tid = db.insert_tailored(mid, "resume text", "keyword", ["playwright"])
            updated = db.update_tailored(tid, "resume text + playwright", ["playwright"])
            self.assertEqual(updated["user_confirmed_skills"], ["playwright"])
            self.assertEqual(updated["added_skills_flagged"], [])
            self.assertIn("playwright", updated["tailored_text"])
        finally:
            db.DB_PATH = old
            tmp.cleanup()


class OfficeAgentsTests(unittest.TestCase):
    """The three new Agent Office specialists."""
    RESUME = ("Priya Nair\nQA Manager | Dallas, TX\n\nSUMMARY\nQA leader.\n\n"
              "TECHNICAL SKILLS\nLanguages: Java, SQL\nTools: Selenium, Jenkins\n")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._old = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "t.db"
        db.init_db()

    def tearDown(self):
        matcher._jd_skills.cache_clear()
        db.DB_PATH = self._old
        self.tmp.cleanup()

    def _job(self, title, desc, employment_type=""):
        return db.insert_job({"source": "t", "source_id": title, "title": title,
                              "company": "Co", "location": "", "remote_flag": 0,
                              "url": "", "description": desc, "posted_at": "",
                              "salary": "", "employment_type": employment_type})

    def _consultant(self, emp_pref=""):
        c = db.create_consultant({"name": "Priya Nair", "emp_pref": emp_pref})
        sk = skills.extract_skills(self.RESUME)
        db.upsert_resume(c["id"], "r.txt", self.RESUME, sk)
        return c

    def test_market_analyst_finds_hot_skills_and_gaps(self):
        self._consultant()
        for i in range(3):
            self._job(f"QA Role {i}", "QA role needing Selenium and Java.")
        for i in range(5):
            self._job(f"Auto Role {i}", "Automation with Playwright and TypeScript.")
        got = office._market_analyst(db.get_settings())
        self.assertEqual(got["window_days"], 30)
        pc = got["consultants"][0]
        hot = {h["skill"]: h["jobs"] for h in pc["hot_skills"]}
        self.assertEqual(hot.get("selenium"), 3)
        gaps = [g["skill"] for g in pc["gaps"]]
        self.assertIn("playwright", gaps)

    def test_outreach_drafts_from_real_data_only(self):
        c = self._consultant()
        jid = self._job("QA Manager", "We need a QA Manager with Selenium and Java.")
        mid = db.insert_match(c["id"], jid, 85.0, {}, [])
        first = office._outreach(db.get_settings())
        self.assertEqual(first["drafted"], 1)
        d = db.list_outreach_drafts()[0]
        self.assertIn("Priya Nair", d["subject"])
        self.assertIn("Selenium", d["body"])          # matched skill, display name
        self.assertNotIn("years", d["body"].lower())  # never invents experience
        self.assertIn("Drafted by BenchPilot Outreach", d["body"])
        second = office._outreach(db.get_settings())  # no duplicates
        self.assertEqual(second["drafted"], 0)
        self.assertEqual(len(db.list_outreach_drafts()), 1)

    def test_compliance_flags_type_mismatch(self):
        c = self._consultant(emp_pref="c2c")
        bad = self._job("QA Manager", "QA Manager role.", employment_type="Full-time")
        good = self._job("QA Engineer", "QA Engineer role.", employment_type="C2C")
        db.insert_match(c["id"], bad, 70.0, {}, [])
        db.insert_match(c["id"], good, 70.0, {}, [])
        got = office._compliance(db.get_settings())
        self.assertEqual(got["issues"], 1)
        item = got["items"][0]
        self.assertEqual(item["issue"], "Employment-type mismatch")
        self.assertIn("C2C", item["detail"])

    def test_autopilot_quarantines_and_scout_hides(self):
        c = self._consultant(emp_pref="c2c")
        jid = self._job("QA Manager", "QA Manager role.", employment_type="Full-time")
        mid = db.insert_match(c["id"], jid, 70.0, {}, [])
        got = office._compliance(db.get_settings())  # autopilot on by default
        self.assertTrue(got["items"][0]["quarantined"])
        self.assertEqual(len(got["actions"]), 1)
        with db.get_conn() as conn:
            flag = conn.execute('SELECT compliance_flag FROM "matches" WHERE id=?',
                                (mid,)).fetchone()["compliance_flag"]
        self.assertEqual(flag, 1)
        # scout no longer surfaces it as a fresh match
        import datetime as _dt
        since = _dt.datetime.now(_dt.timezone.utc) - _dt.timedelta(hours=25)
        fresh = office._scout(db.get_settings(), since)["new_matches"]
        self.assertEqual(fresh, 0)
        # restore works
        db.set_compliance_flag(mid, 0)
        fresh = office._scout(db.get_settings(), since)["new_matches"]
        self.assertEqual(fresh, 1)

    def test_autopilot_off_only_reports(self):
        db.set_setting("office_autopilot", "0")
        c = self._consultant(emp_pref="c2c")
        jid = self._job("QA Manager", "QA Manager role.", employment_type="Full-time")
        mid = db.insert_match(c["id"], jid, 70.0, {}, [])
        got = office._compliance(db.get_settings())
        self.assertEqual(got["issues"], 1)
        self.assertFalse(got["items"][0]["quarantined"])
        self.assertEqual(got["actions"], [])
        with db.get_conn() as conn:
            flag = conn.execute('SELECT compliance_flag FROM "matches" WHERE id=?',
                                (mid,)).fetchone()["compliance_flag"]
        self.assertEqual(flag, 0)

    def test_market_analyst_expands_queries_for_hot_gaps(self):
        self._consultant()  # skills: java, sql, selenium, jenkins - no playwright
        for i in range(4):
            self._job(f"Playwright QA {i}", "Senior QA Automation with Playwright.")
        got = office._market_analyst(db.get_settings())
        self.assertTrue(got["actions"], "expected a query-expansion action")
        queries = json.loads(db.get_setting("office_expanded_queries"))
        titles = [q["title"].lower() for q in queries]
        self.assertTrue(any("playwright" in t for t in titles))
        # the typed queries are left alone
        self.assertEqual(json.loads(db.get_setting("search_queries")),
                         db.DEFAULT_SEARCH_QUERIES)
        # the collector really searches for it, even with many consultants
        db.set_setting("max_queries", "3")
        qs = autolearn.queries_for_run(db.get_settings())
        self.assertTrue(any("playwright" in q["title"].lower() for q in qs))
        self.assertLessEqual(len(qs), 3)
        # second run must not duplicate the query
        n = len(queries)
        office._market_analyst(db.get_settings())
        self.assertEqual(len(json.loads(db.get_setting("office_expanded_queries"))), n)

    def test_watchdog_leaves_followup_note_once(self):
        c = self._consultant()
        jid = self._job("QA Manager", "QA Manager role.")
        mid = db.insert_match(c["id"], jid, 70.0, {}, [])
        aid = db.insert_application(mid, None) if hasattr(db, "insert_application") else None
        if aid is None:
            # fall back to direct insert matching the schema
            with db.get_conn() as conn:
                cur = conn.execute(
                    'INSERT INTO applications(match_id, status, updated_at) VALUES (?, ?, ?)',
                    (mid, "queued", "2020-01-01T00:00:00+00:00"))
                conn.commit()
                aid = cur.lastrowid
        got = office._watchdog(db.get_settings())
        self.assertEqual(got["stale"], 1)
        self.assertEqual(len(got["actions"]), 1)
        with db.get_conn() as conn:
            notes = conn.execute("SELECT notes FROM applications WHERE id=?",
                                 (aid,)).fetchone()["notes"]
        self.assertIn("[Watchdog]", notes)
        # second run: no duplicate note, no duplicate action
        got2 = office._watchdog(db.get_settings())
        self.assertEqual(got2["actions"], [])
        with db.get_conn() as conn:
            notes2 = conn.execute("SELECT notes FROM applications WHERE id=?",
                                  (aid,)).fetchone()["notes"]
        self.assertEqual(notes2.count("[Watchdog]"), 1)

    def test_run_office_includes_all_six_agents(self):
        self._consultant()
        self._job("QA Manager", "QA Manager with Selenium.")
        br = office.run_office()
        self.assertEqual(set(br["agents"].keys()),
                         {"scout", "tailor", "watchdog",
                          "market_analyst", "outreach", "compliance"})

    # ---- fixes added on top of Autopilot ----
    def _match(self, cid, title, score, missing=(), company="Co"):
        jid = db.insert_job({"source": "t", "source_id": title, "title": title,
                             "company": company, "location": "", "remote_flag": 0,
                             "url": "", "description": "Requirements: Java, Selenium, Kafka",
                             "posted_at": "", "salary": "", "employment_type": ""})
        return db.insert_match(cid, jid, score, {}, list(missing))

    def test_tailor_cutoff_setting_and_potential_score(self):
        c = self._consultant()
        self._match(c["id"], "A", 66.0, ["kafka"])
        self._match(c["id"], "B", 50.0)
        got = office._tailor(db.get_settings())
        self.assertEqual((got["cutoff"], got["drafted"]), (65, 1))
        d = got["drafts"][0]
        self.assertEqual(d["missing"], ["kafka"])
        self.assertGreater(d["potential_score"], d["score"])
        db.set_setting("tailor_cutoff", "40")
        db.set_setting("tailor_max_per_run", "1")
        self.assertEqual(office._tailor(db.get_settings())["drafted"], 1)  # cap holds

    def test_quarantined_matches_get_no_drafts(self):
        c = self._consultant()
        mid = self._match(c["id"], "A", 90.0)
        db.set_compliance_flag(mid, 1)
        self.assertEqual(office._tailor(db.get_settings())["drafted"], 0)
        self.assertEqual(office._outreach(db.get_settings())["drafted"], 0)

    def test_compliance_runs_before_drafting(self):
        c = db.create_consultant({"name": "Priya Nair", "emp_pref": "c2c"})
        db.upsert_resume(c["id"], "r.txt", self.RESUME, skills.extract_skills(self.RESUME))
        jid = db.insert_job({"source": "t", "source_id": "w2", "title": "QA W2",
                             "company": "Co", "location": "", "remote_flag": 0,
                             "url": "", "description": "W2 only. Selenium Java",
                             "posted_at": "", "salary": "", "employment_type": ""})
        mid = db.insert_match(c["id"], jid, 90.0, {}, [])
        with db.get_conn() as conn:
            conn.execute("UPDATE jobs SET emp_tags=',w2,' WHERE id=?", (jid,))
            conn.commit()
        br = office.run_office()["agents"]
        self.assertEqual(br["compliance"]["issues"], 1)
        self.assertEqual(br["tailor"]["drafted"], 0)      # quarantined first
        self.assertEqual(br["outreach"]["drafted"], 0)

    def test_scout_looks_back_past_a_sleep_gap(self):
        c = self._consultant()
        self._match(c["id"], "A", 70.0)
        with db.get_conn() as conn:
            conn.execute('UPDATE "matches" SET created_at=?', ("2020-01-02T00:00:00+00:00",))
            conn.commit()
        self.assertEqual(office.run_office()["agents"]["scout"]["new_matches"], 0)
        db.set_setting("office_last_run_at", "2020-01-01T00:00:00+00:00")
        self.assertEqual(office.run_office()["agents"]["scout"]["new_matches"], 1)

    def test_run_and_store_marks_day_done(self):
        self._consultant()
        office.run_and_store()
        self.assertTrue(db.get_setting("office_last_run_date"))
        self.assertTrue(db.get_setting("office_last_run_at"))

    def test_watchdog_bad_date_is_not_stale(self):
        c = self._consultant()
        mid = self._match(c["id"], "A", 70.0)
        aid = db.queue_application(mid)["id"]
        with db.get_conn() as conn:
            conn.execute("UPDATE applications SET updated_at='garbage' WHERE id=?", (aid,))
            conn.commit()
        self.assertEqual(office._watchdog(db.get_settings())["stale"], 0)



class ContactTests(unittest.TestCase):
    def test_header_is_read(self):
        c = contact.parse_contact(
            "SHARAN MURALI\nSenior QA Test Manager | Test Lead\n"
            "sharmurali9@gmail.com | +1 (714) 860-2332\nSUMMARY\nx", "a.docx")
        self.assertEqual((c["name"], c["email"], c["phone"]),
                         ("Sharan Murali", "sharmurali9@gmail.com", "+1 (714) 860-2332"))
        c = contact.parse_contact(
            "Ravi Kumar\nJava Full Stack Developer\n"
            "Dallas, TX | ravi.k@mail.com | (469) 555-0123", "r.pdf")
        self.assertEqual((c["name"], c["location"]), ("Ravi Kumar", "Dallas, TX"))

    def test_name_falls_back_to_file_name(self):
        c = contact.parse_contact("RESUME\nSenior Java Developer\nJava", "Priya_Sharma_CV_final.pdf")
        self.assertEqual(c["name"], "Priya Sharma")
        self.assertEqual(contact.parse_contact("", "")["name"], "New consultant")

    def test_job_title_is_not_a_name(self):
        c = contact.parse_contact("Senior Java Developer\nJane Doe\n", "x.pdf")
        self.assertEqual(c["name"], "Jane Doe")


class ResumeImportTests(unittest.TestCase):
    """Google Drive folder import (network calls are faked)."""
    def setUp(self):
        from app import resume_import
        self.ri = resume_import
        self.tmp = tempfile.TemporaryDirectory()
        self._old = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "i.db"
        db.init_db()
        self.files = [
            {"id": "f1", "name": "Sharan Murali Resume.docx", "mimeType": self.ri.DOCX,
             "md5Checksum": "aaa"},
            {"id": "f2", "name": "Ravi.pdf", "mimeType": "application/pdf",
             "md5Checksum": "bbb"},
            {"id": "f3", "name": "photo.png", "mimeType": "image/png", "md5Checksum": "c"},
            {"id": "f4", "name": "Priya (Google Doc)", "mimeType": self.ri.GDOC,
             "modifiedTime": "2026-10-01T00:00:00Z"},
        ]
        self._lf, self._dl = self.ri.list_folder, self.ri.download
        self.ri.list_folder = lambda fid, key: list(self.files)
        self.ri.download = lambda f, key: b"data-" + f["id"].encode()
        self.settings = {"resume_folder_url":
                         "https://drive.google.com/drive/folders/1AbCdEfGhIjKlMnOp?usp=sharing",
                         "google_api_key": "k"}
        self.seen = []

    def tearDown(self):
        self.ri.list_folder, self.ri.download = self._lf, self._dl
        db.DB_PATH = self._old
        self.tmp.cleanup()

    def ingest(self, name, data, kick):
        self.seen.append((name, kick))
        if name.startswith("Ravi"):
            raise ValueError("could not read this file")
        return {"consultant_id": len(self.seen), "name": name.split(".")[0],
                "created": True, "new_matches": 2}

    def test_folder_id_from_links(self):
        f = self.ri.folder_id
        self.assertEqual(f("https://drive.google.com/drive/folders/1AbCdEfGhIjKlMnOp?usp=sharing"),
                         "1AbCdEfGhIjKlMnOp")
        self.assertEqual(f("https://drive.google.com/open?id=1AbCdEfGhIjKlMnOp"), "1AbCdEfGhIjKlMnOp")
        self.assertEqual(f("1AbCdEfGhIjKlMnOp"), "1AbCdEfGhIjKlMnOp")
        self.assertEqual(f("not a link"), "")

    def test_reads_resumes_once_and_skips_other_files(self):
        r = self.ri.run(self.ingest, self.settings)
        self.assertEqual(len(r["added"]), 2)          # f1 + the Google Doc
        self.assertEqual(len(r["failed"]), 1)         # Ravi.pdf could not be read
        self.assertEqual(r["skipped"], 1)             # the .png
        names = [n for n, _ in self.seen]
        self.assertIn("Priya (Google Doc).docx", names)   # Google Doc exported as .docx
        self.assertEqual([k for _, k in self.seen].count(True), 1)  # job search kicked once
        self.seen.clear()
        r2 = self.ri.run(self.ingest, self.settings)  # nothing new -> nothing read again
        self.assertEqual((r2["added"], r2["failed"], self.seen), ([], [], []))

    def test_changed_file_is_read_again(self):
        self.ri.run(self.ingest, self.settings)
        self.seen.clear()
        self.files[0]["md5Checksum"] = "new"
        r = self.ri.run(self.ingest, self.settings)
        self.assertEqual([n for n, _ in self.seen], ["Sharan Murali Resume.docx"])

    def test_not_set_up_and_drive_errors_do_not_raise(self):
        self.assertIn("not set up", self.ri.run(self.ingest, {})["error"])
        def boom(fid, key):
            raise self.ri.DriveError("folder not found")
        self.ri.list_folder = boom
        self.assertEqual(self.ri.run(self.ingest, self.settings)["error"], "folder not found")

    def test_download_problem_is_retried_next_time(self):
        calls = {"n": 0}
        def flaky(f, key):
            calls["n"] += 1
            if calls["n"] == 1:
                raise self.ri.DriveError("network")
            return b"x"
        self.ri.download = flaky
        self.files = self.files[:1]
        r = self.ri.run(self.ingest, self.settings)
        self.assertEqual(len(r["failed"]), 1)
        r = self.ri.run(self.ingest, self.settings)
        self.assertEqual(len(r["added"]), 1)

    def test_per_run_limit(self):
        r = self.ri.run(self.ingest, self.settings, files_per_run=1)
        self.assertEqual(r["waiting"], 2)


class ResumeAiTests(unittest.TestCase):
    """AI-checked resume reading (the AI call is faked)."""
    TEXT = ("SHARAN MURALI\nSenior QA Test Manager | Test Lead\n"
            "sharmurali9@gmail.com | +1 (714) 860-2332 | Irvine, CA\n\nSUMMARY\nQA leader.\n")
    KEY = {"llm_base_url": "http://x", "llm_api_key": "k", "llm_model": "m"}

    def setUp(self):
        from app import resume_ai
        self.ra = resume_ai
        self._orig = resume_ai.ai_extract

    def tearDown(self):
        self.ra.ai_extract = self._orig

    def fake(self, **kw):
        base = {"name": "Sharan Murali", "email": "sharmurali9@gmail.com",
                "phone": "+1 (714) 860-2332", "location": "Irvine, CA",
                "job_title": "Senior QA Test Manager", "confidence": "high"}
        base.update(kw)
        self.ra.ai_extract = lambda text, settings: base

    def test_good_answer_needs_no_check(self):
        self.fake()
        r = self.ra.read_contact(self.TEXT, "Sharan.docx", self.KEY)
        self.assertFalse(r["needs_check"])
        self.assertEqual(r["info"]["name"], "Sharan Murali")
        self.assertTrue(r["used_ai"])

    def test_invented_email_is_thrown_away(self):
        self.fake(email="sharan@invented.com")
        r = self.ra.read_contact(self.TEXT, "Sharan.docx", self.KEY)
        self.assertEqual(r["info"]["email"], "sharmurali9@gmail.com")   # kept the real one

    def test_invented_name_is_thrown_away(self):
        self.fake(name="John Smith")
        r = self.ra.read_contact(self.TEXT, "Sharan.docx", self.KEY)
        self.assertEqual(r["info"]["name"], "Sharan Murali")

    def test_low_confidence_is_flagged(self):
        self.fake(confidence="low")
        r = self.ra.read_contact(self.TEXT, "Sharan.docx", self.KEY)
        self.assertTrue(r["needs_check"])
        self.assertTrue(any("not sure" in n for n in r["notes"]))

    def test_name_disagreement_is_flagged(self):
        self.fake(name="Sharan Murali Test")      # also in the text start? no: squash differs
        r = self.ra.read_contact(self.TEXT, "Sharan.docx", self.KEY)
        self.assertEqual(r["info"]["name"], "Sharan Murali")     # unsupported -> rules value

    def test_one_word_name_above_contact_line(self):
        self.ra.ai_extract = lambda text, settings: None
        t = ("RAJESWARI\n281-627-6787 | rajisplmtc@gmail.com\nPROFESSIONAL SUMMARY\n"
             "Teamcenter PLM professional.\nApplication Support\nDevelopment\n")
        r = self.ra.read_contact(t, "Rajeswari_TC_PLM_Resume.pdf", {})
        self.assertEqual(r["info"]["name"], "Rajeswari")      # not "Application Support"
        self.assertTrue(r["needs_check"])
        self.assertTrue(any("only one name" in n for n in r["notes"]))

    def test_no_city_is_not_a_mistake(self):
        self.ra.ai_extract = lambda text, settings: None
        t = "SHARAN MURALI\nQA Manager\nsharmurali9@gmail.com | +1 (714) 860-2332\n"
        self.assertFalse(self.ra.read_contact(t, "x.docx", {})["needs_check"])

    def test_without_ai_rules_are_checked_and_flagged_when_weak(self):
        self.ra.ai_extract = lambda text, settings: None
        ok = self.ra.read_contact(self.TEXT, "Sharan.docx", {})
        self.assertFalse(ok["needs_check"])
        self.assertFalse(ok["used_ai"])
        weak = self.ra.read_contact("Java developer\nskills: java", "x.pdf", {})
        self.assertTrue(weak["needs_check"])
        self.assertIn("no valid email found", weak["notes"])

    def test_bad_ai_call_falls_back_to_rules(self):
        self.ra.ai_extract = self._orig
        import requests as rq
        old = rq.post
        rq.post = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("down"))
        try:
            r = self.ra.read_contact(self.TEXT, "Sharan.docx", self.KEY)
        finally:
            rq.post = old
        self.assertFalse(r["used_ai"])
        self.assertEqual(r["info"]["email"], "sharmurali9@gmail.com")

    def test_json_in_code_fence_is_read(self):
        self.ra.ai_extract = self._orig
        import requests as rq
        class R:
            def raise_for_status(self): pass
            def json(self):
                return {"choices": [{"message": {"content":
                    '```json\n{"name":"Sharan Murali","email":"sharmurali9@gmail.com",'
                    '"phone":"","location":"Irvine, CA","job_title":"QA","confidence":"high"}\n```'}}]}
        old = rq.post
        rq.post = lambda *a, **k: R()
        try:
            got = self.ra.ai_extract(self.TEXT, self.KEY)
        finally:
            rq.post = old
        self.assertEqual(got["email"], "sharmurali9@gmail.com")

    def test_flag_is_stored_and_cleared(self):
        tmp = tempfile.TemporaryDirectory()
        old = db.DB_PATH
        db.DB_PATH = Path(tmp.name) / "c.db"
        try:
            db.init_db()
            c = db.create_consultant({"name": "A B"})
            db.set_consultant_check(c["id"], True, ["no valid email found", "no city found"])
            row = db.get_consultant(c["id"])
            self.assertEqual((row["needs_check"], row["check_notes"]),
                             (1, "no valid email found; no city found"))
            db.set_consultant_check(c["id"], False, "")
            self.assertEqual(db.get_consultant(c["id"])["needs_check"], 0)
        finally:
            db.DB_PATH = old
            tmp.cleanup()


class JenkinAliasTests(unittest.TestCase):
    def test_typo_maps_to_jenkins_and_stale_learned_skill_is_dropped(self):
        tmp = tempfile.TemporaryDirectory()
        old = db.DB_PATH
        db.DB_PATH = Path(tmp.name) / "j.db"
        try:
            db.init_db()
            self.assertIn("jenkins", skills.extract_skills("Configuration Tools SVN, Git, Jenkin"))
            self.assertNotIn("jenkin", skills.extract_skills("Configuration Tools SVN, Git, Jenkin"))
            db.add_learned_skills(["jenkin"], source="old")      # learned before the alias existed
            autolearn.maintain()
            self.assertNotIn("jenkin", [r["skill"] for r in db.list_learned_skills()])
            c = db.create_consultant({"name": "R"})
            db.upsert_resume(c["id"], "r.txt", "R\nSkills\nSVN, Git, Jenkin", ["jenkin", "git"])
            autolearn.maintain()
            self.assertIn("jenkins", db.get_resume(c["id"])["skills_json"])
        finally:
            skills.set_learned([])
            db.DB_PATH = old
            tmp.cleanup()


class RoleFromSummaryTests(unittest.TestCase):
    RAJI = ("RAJESWARI\n281-627-6787 | rajisplmtc@gmail.com\nPROFESSIONAL SUMMARY\n"
            "Results-driven Teamcenter PLM Professional with 8+ years of experience in "
            "Teamcenter administration.\nAREAS OF EXPERTISE\n")

    def test_role_is_read_from_the_summary_when_there_is_no_title_line(self):
        self.assertEqual(matcher.role_from_summary(self.RAJI), "Teamcenter PLM")
        self.assertEqual(matcher.consultant_title(self.RAJI), "Teamcenter PLM")
        self.assertEqual(autolearn.headline_title(self.RAJI, ["agile", "c++"]), "Teamcenter PLM")

    def test_title_line_still_wins(self):
        t = "SHARAN MURALI\nSenior QA Test Manager | Test Lead\nsharmurali9@gmail.com\n"
        self.assertEqual(matcher.consultant_title(t), "Senior QA Test Manager")
        self.assertEqual(autolearn.headline_title(t, []), "Senior QA Test Manager")

    def test_nothing_sensible_gives_empty(self):
        self.assertEqual(matcher.role_from_summary("hello world\nsome text"), "")


class FakeStore:
    def __init__(self):
        self.files = {}
    def exists(self, name):
        return name in self.files
    def put(self, name, path):
        self.files[name] = open(path, "rb").read()
    def get(self, name, path):
        open(path, "wb").write(self.files[name])


class PersistTests(unittest.TestCase):
    def setUp(self):
        from app import persist
        self.persist = persist
        self.tmp = tempfile.TemporaryDirectory()
        self.old = db.DB_PATH
        self.old_sig = persist._last_sig
        persist._last_sig = {}
        db.DB_PATH = Path(self.tmp.name) / "a" / "b.db"
        self.store = FakeStore()

    def tearDown(self):
        db.DB_PATH = self.old
        self.persist._last_sig = self.old_sig
        self.tmp.cleanup()

    def test_backup_then_restore_after_wipe(self):
        db.init_db()
        db.set_setting("llm_api_key", "gsk_test")
        c = db.create_consultant({"name": "Sharan Murali", "email": "s@x.com"})
        self.assertTrue(self.persist.backup(self.store))
        self.assertIn(self.persist.OBJECT, self.store.files)
        db.DB_PATH.unlink()                      # Republish wipes the disk
        self.assertTrue(self.persist.restore(self.store))
        db.init_db()
        self.assertEqual(db.get_setting("llm_api_key"), "gsk_test")
        self.assertEqual(db.get_consultant(c["id"])["name"], "Sharan Murali")

    def test_unchanged_database_is_not_uploaded_twice(self):
        db.init_db()
        self.assertTrue(self.persist.backup(self.store))
        self.assertFalse(self.persist.backup(self.store))
        self.assertTrue(self.persist.backup(self.store, force=True))

    def test_newer_local_database_is_not_overwritten(self):
        db.init_db()
        db.create_consultant({"name": "Old One", "email": "o@x.com"})
        self.persist.backup(self.store)
        db.create_consultant({"name": "New Two", "email": "n@x.com"})      # newer than the backup
        db.set_setting("backup_stamp", "2999-01-01T00:00:00+00:00")
        self.assertFalse(self.persist.restore(self.store))
        names = [c["name"] for c in db.list_consultants()]
        self.assertIn("New Two", names)

    def test_nothing_in_storage_does_nothing(self):
        self.assertFalse(self.persist.restore(self.store))

    def test_failure_never_raises(self):
        class Broken(FakeStore):
            def put(self, name, path):
                raise RuntimeError("no bucket")
        db.init_db()
        self.assertFalse(self.persist.backup(Broken()))
        self.assertIn("backup failed", self.persist.status()["last_error"])


class EnvSettingsTests(unittest.TestCase):
    def test_secret_fills_empty_key_but_saved_value_wins(self):
        tmp = tempfile.TemporaryDirectory()
        old = db.DB_PATH
        db.DB_PATH = Path(tmp.name) / "e.db"
        env = dict(os.environ)
        try:
            db.init_db()
            os.environ["LLM_API_KEY"] = "from-secret"
            os.environ["LLM_MODEL"] = "model-secret"
            self.assertEqual(db.get_setting("llm_api_key"), "from-secret")
            db.set_setting("llm_model", "typed-in-settings")
            self.assertEqual(db.get_setting("llm_model"), "typed-in-settings")
        finally:
            os.environ.clear(); os.environ.update(env)
            db.DB_PATH = old
            tmp.cleanup()


class DiceConnectorTests(unittest.TestCase):
    SCHEMA = {"type": "object", "required": ["keyword"], "properties": {
        "keyword": {"type": "string"},
        "location": {"type": "string"},
        "workplace_types": {"type": "array", "items": {"enum": ["Remote", "On-Site", "Hybrid"]}},
        "employment_types": {"type": "array", "items": {"enum": ["FULLTIME", "PARTTIME", "CONTRACTS", "THIRD_PARTY"]}},
        "posted_date": {"type": "string", "enum": ["ONE", "THREE", "SEVEN"]},
        "page_size": {"type": "integer", "maximum": 100}}}

    def test_arguments_follow_the_tools_own_schema(self):
        from app.sources import dice
        a = dice.build_args(self.SCHEMA, {"title": "QA Automation", "location": "Remote"},
                            {"collect_emp_types": "c2c,w2,fulltime"})
        self.assertEqual(a["keyword"], "QA Automation")
        self.assertNotIn("location", a)
        self.assertEqual(a["workplace_types"], ["Remote"])
        self.assertEqual(set(a["employment_types"]), {"FULLTIME", "CONTRACTS", "THIRD_PARTY"})
        self.assertEqual(a["posted_date"], "SEVEN")
        self.assertEqual(a["page_size"], 50)
        b = dice.build_args(self.SCHEMA, {"title": "QA", "location": "Dallas, TX"}, {})
        self.assertEqual(b["location"], "Dallas, TX")
        self.assertNotIn("workplace_types", b)

    def test_other_parameter_names_still_work(self):
        from app.sources import dice
        a = dice.build_args({"properties": {"query": {"type": "string"},
                                            "remote": {"type": "boolean"}}},
                            {"title": "SDET", "location": ""}, {})
        self.assertEqual(a, {"query": "SDET", "remote": True})

    def test_answer_is_turned_into_jobs(self):
        from app.sources import dice
        payload = {"data": [{"id": "abc-123", "title": "QA Automation Engineer",
                             "companyName": "Acme", "detailsPageUrl": "https://www.dice.com/job-detail/abc-123",
                             "jobLocation": {"displayName": "Austin, TX"},
                             "workplaceTypes": ["Remote"], "employmentType": "Contract",
                             "summary": "Selenium, Java", "postedDate": "2026-10-20"},
                            {"id": "x", "title": ""}], "meta": {"total": 2}}
        res = {"content": [{"type": "text", "text": json.dumps(payload)}]}
        jobs = dice.parse_result(res)
        self.assertEqual(len(jobs), 1)
        j = jobs[0]
        self.assertEqual((j["source"], j["company"], j["location"]), ("dice", "Acme", "Austin, TX"))
        self.assertTrue(j["remote_flag"])
        self.assertTrue(j["url"].startswith("https://www.dice.com/"))
        self.assertEqual(dice.parse_result({"structuredContent": {"jobs": [
            {"jobTitle": "SDET", "url": "/job-detail/9"}]}})[0]["url"],
            "https://www.dice.com/job-detail/9")

    def test_unknown_answer_shows_a_useful_error(self):
        from app.sources import dice
        with self.assertRaises(RuntimeError) as cm:
            dice.parse_result({"content": [{"type": "text", "text": "Here are some jobs: none"}]})
        self.assertIn("unexpected format", str(cm.exception))

    def test_event_stream_answers_are_read(self):
        from app.sources import dice
        txt = 'event: message\ndata: {"jsonrpc":"2.0","id":2,"result":{"tools":[]}}\n\n'
        self.assertEqual(dice._parse_rpc(txt, "text/event-stream", 2)["result"], {"tools": []})

    def test_full_conversation_with_a_fake_dice(self):
        from app.sources import dice
        schema = self.SCHEMA

        class R:
            def __init__(self, body, ctype="application/json"):
                self.text, self.status_code = json.dumps(body), 200
                self.headers = {"Content-Type": ctype, "Mcp-Session-Id": "s1"}
            def raise_for_status(self): pass

        sent = []

        def post(url, data=None, timeout=None):
            msg = json.loads(data)
            sent.append(msg)
            m = msg.get("method")
            if m == "initialize":
                return R({"jsonrpc": "2.0", "id": msg["id"], "result": {"capabilities": {}}})
            if m == "tools/list":
                return R({"jsonrpc": "2.0", "id": msg["id"], "result": {"tools": [
                    {"name": "get_company"}, {"name": "search_jobs", "inputSchema": schema}]}})
            if m == "tools/call":
                return R({"jsonrpc": "2.0", "id": msg["id"], "result": {"content": [
                    {"type": "text", "text": json.dumps({"data": [
                        {"id": "1", "title": "SDET", "companyName": "B",
                         "detailsPageUrl": "https://www.dice.com/job-detail/1"}]})}]}})
            return R({})
        old = dice.requests.Session
        fake_post = post

        class FakeSession:
            def __init__(self):
                self.headers = {}

            def post(self, url, data=None, timeout=None):
                return fake_post(url, data=data, timeout=timeout)
        dice.requests.Session = FakeSession
        try:
            jobs = dice.fetch({"title": "SDET", "location": "Remote"}, {})
        finally:
            dice.requests.Session = old
        self.assertEqual([j["title"] for j in jobs], ["SDET"])
        call = [m for m in sent if m.get("method") == "tools/call"][0]
        self.assertEqual(call["params"]["name"], "search_jobs")
        self.assertEqual(call["params"]["arguments"]["keyword"], "SDET")

    def test_on_by_default_and_can_be_turned_off(self):
        from app.sources import dice
        self.assertTrue(dice.enabled({}))
        self.assertFalse(dice.enabled({"dice_enabled": "0"}))


class LiveSearchRelevanceTests(unittest.TestCase):
    def test_only_jobs_about_the_search_are_shown(self):
        from app.relevance import title_fits
        q = "QA Automation Engineer"
        for ok in ("QA Automation Engineer", "Senior Test Automation Engineer",
                   "SDET - Automation", "Quality Assurance Automation Lead"):
            self.assertTrue(title_fits(q, ok), ok)
        for bad in ("Frontend Engineer", "QA Supervisor", "Inside Sales Contractor"):
            self.assertFalse(title_fits(q, bad), bad)
        self.assertTrue(title_fits("Java Developer", "Sr. Java Developer (C2C)"))
        self.assertFalse(title_fits("Java Developer", "Python Developer"))


class JSearchBudgetTests(unittest.TestCase):
    def test_daily_and_monthly_limits(self):
        import collector
        tmp = tempfile.TemporaryDirectory()
        old = db.DB_PATH
        db.DB_PATH = Path(tmp.name) / "j.db"
        try:
            db.init_db()
            s = db.get_settings()
            self.assertEqual(collector.jsearch_room(s), (6, 195))
            db.usage_add("jsearch", 4)
            self.assertEqual(collector.jsearch_room(s)[0], 2)
            db.usage_add("jsearch", 190)
            self.assertEqual(collector.jsearch_room(s), (0, 1))
            self.assertEqual(db.usage_month("jsearch"), 194)
        finally:
            db.DB_PATH = old
            tmp.cleanup()


class JSearchAnswerTests(unittest.TestCase):
    ROW = {"job_id": "a1", "job_title": "QA Automation Engineer", "employer_name": "Acme",
           "job_city": "Austin", "job_state": "TX", "job_is_remote": True,
           "job_apply_link": "https://x/apply", "job_publisher": "LinkedIn",
           "job_employment_type": "CONTRACTOR", "job_description": "Selenium"}

    def test_both_answer_shapes_work(self):
        from app.sources import jsearch
        self.assertEqual(len(jsearch.job_rows({"data": [self.ROW]})), 1)
        self.assertEqual(len(jsearch.job_rows({"status": "OK", "data": {"jobs": [self.ROW],
                                                                         "cursor": "x"}})), 1)
        self.assertEqual(jsearch.job_rows({"data": "nothing"}), [])

    def test_job_shows_where_it_was_posted(self):
        from app.sources import jsearch
        j = jsearch.to_job(self.ROW)
        self.assertEqual((j["source"], j["location"], j["company"]),
                         ("linkedin", "Austin, TX", "Acme"))
        self.assertEqual(jsearch.to_job({**self.ROW, "job_publisher": "Indeed"})["source"], "indeed")
        self.assertEqual(jsearch.to_job({**self.ROW, "job_publisher": "Jooble"})["source"], "jooble")
        self.assertEqual(jsearch.to_job({**self.ROW, "job_publisher": "Talent.com"})["source"], "talent.com")
        self.assertEqual(jsearch.to_job({**self.ROW, "job_publisher": ""})["source"], "jsearch")


class MonsterTests(unittest.TestCase):
    def test_input_and_output(self):
        from app.sources import monster
        b = monster.build_input({"title": "QA Automation", "location": "Remote"}, {})
        self.assertEqual(b, {"query": "QA Automation", "country": "US", "maxResults": 30,
                             "location": "Remote"})
        self.assertNotIn("location", monster.build_input({"title": "QA", "location": "USA"}, {}))
        j = monster.to_job({"jobId": "m1", "title": "SDET", "company": "Acme",
                            "city": "Dallas", "state": "TX", "isRemote": False,
                            "url": "https://www.monster.com/job-openings/m1",
                            "salaryMin": 50, "salaryMax": 60, "postedDate": "2026-10-20"})
        self.assertEqual((j["source"], j["location"], j["salary"]), ("monster", "Dallas, TX", "$50 - $60"))
        self.assertIsNone(monster.to_job({"title": "x"}))

    def test_collector_runs_monster_then_waits(self):
        import collector
        from app.sources import monster
        tmp = tempfile.TemporaryDirectory()
        old_db, old_fetch = db.DB_PATH, monster.fetch
        db.DB_PATH = Path(tmp.name) / "mo.db"
        calls = []

        def fake(q, settings):
            calls.append(q["title"])
            return [{"source": "monster", "source_id": "1", "title": "SDET", "company": "A",
                     "location": "Remote", "remote_flag": True, "url": "https://m/1",
                     "description": "selenium", "posted_at": "", "salary": "",
                     "employment_type": ""}]
        try:
            db.init_db()
            db.set_setting("apify_token", "apify_api_test")
            db.set_setting("enabled_sources", "[]")
            db.set_setting("dice_enabled", "0")
            db.set_setting("auto_queries", "0")
            db.set_setting("search_queries", json.dumps([{"title": "SDET", "location": "Remote"}]))
            monster.fetch = fake
            s1 = collector.run()
            self.assertEqual(s1["sources"]["monster"]["new"], 1)
            s2 = collector.run()
            self.assertIn("waiting", s2["sources"]["monster"]["status"])
            self.assertEqual(calls, ["SDET"])
        finally:
            monster.fetch = old_fetch
            db.DB_PATH = old_db
            tmp.cleanup()


class TargetRoleTests(unittest.TestCase):
    TEXT = ("Jane D\n(555)123-4567\njane@x.com\nSummary\nAround 10 years of experience, "
            "with specialization in Software Quality Assurance Testing, API testing.\n")
    DESC = "selenium java jenkins docker aws playwright api testing"

    def _c(self, **kw):
        c = {"name": "Jane D", "raw_text": self.TEXT,
             "filename": "Jane_Software_Development_Engineer_in_Test_SDET.docx",
             "skills": ["selenium", "java", "jenkins", "docker", "aws", "playwright"]}
        c.update(kw)
        return c

    def test_name_is_not_a_title_and_roles_are_found(self):
        titles = matcher.titles_for(self._c())
        self.assertNotIn("Jane D", titles)
        self.assertIn("Software Development Engineer in Test SDET", titles)
        self.assertTrue(matcher.has_role(self._c()))

    def test_sdet_gets_qa_jobs_but_not_devops_or_java(self):
        c = self._c()
        def fits(title):
            r = matcher.score(c, {"title": title, "description": self.DESC,
                                  "location": "Remote", "remote_flag": True})
            return matcher.fits_role(r["breakdown"], matcher.has_role(c))
        for ok in ("QA Automation Engineer", "SDET", "Senior Test Automation Engineer",
                   "Software Development Engineer in Test"):
            self.assertTrue(fits(ok), ok)
        for bad in ("Senior DevOps Engineer", "Lead Java Developer", "DEVOPS ENGINEER"):
            self.assertFalse(fits(bad), bad)

    def test_typed_target_roles_win_for_matching_and_searching(self):
        c = self._c(target_roles="SDET, QA Automation Engineer")
        self.assertEqual(matcher.titles_for(c), ["SDET", "QA Automation Engineer"])
        qs = autolearn.build_queries([c], [], 8)
        self.assertEqual([q["title"] for q in qs], ["SDET", "QA Automation Engineer"])

    def test_unknown_role_still_allows_a_perfect_skill_fit(self):
        self.assertTrue(matcher.fits_role({"title": 0, "skill": 95}, strict=False))
        self.assertFalse(matcher.fits_role({"title": 0, "skill": 95}, strict=True))

    def test_target_roles_saved_on_consultant(self):
        tmp = tempfile.TemporaryDirectory()
        old = db.DB_PATH
        db.DB_PATH = Path(tmp.name) / "tr.db"
        try:
            db.init_db()
            c = db.create_consultant({"name": "A", "target_roles": "SDET"})
            self.assertEqual(c["target_roles"], "SDET")
            self.assertEqual(db.update_consultant(c["id"], {"target_roles": "QA"})["target_roles"], "QA")
        finally:
            db.DB_PATH = old
            tmp.cleanup()


class AgentFixTests(unittest.TestCase):
    def test_foreign_region_with_us_suffix_is_not_usa(self):
        from app import usa
        self.assertFalse(usa.is_us_job({"location": "Medellin, AN, US"}))
        self.assertTrue(usa.is_us_job({"location": "Dallas, TX, US"}))
        self.assertTrue(usa.is_us_job({"location": "US"}))

    def test_same_job_in_many_cities_is_one_match(self):
        tmp = tempfile.TemporaryDirectory()
        old = db.DB_PATH
        db.DB_PATH = Path(tmp.name) / "dup.db"
        try:
            db.init_db()
            c = db.create_consultant({"name": "Q", "target_roles": "QA Automation Engineer"})
            db.upsert_resume(c["id"], "q.txt", "Q\nQA Automation Engineer\nselenium java playwright",
                             ["selenium", "java", "playwright"])
            for i, city in enumerate(["Boston, MA", "Phoenix, AZ", "Dallas, TX"]):
                db.insert_job({"source": "adzuna", "source_id": str(i),
                               "title": "Senior QA Automation Engineer - Playwright & Selenium",
                               "company": "Cognizant", "location": city, "url": "u",
                               "description": "selenium java playwright", "posted_at": ""})
            matcher.run_all(threshold=0)
            self.assertEqual(len(db.list_matches()), 1)
        finally:
            db.DB_PATH = old
            tmp.cleanup()

    def test_outreach_email_has_no_internal_score(self):
        src = open(Path(__file__).resolve().parent.parent / "app" / "office.py").read()
        self.assertNotIn("internal match score", src)


class RoleFitTests(unittest.TestCase):
    def test_fits_role_rule(self):
        self.assertTrue(matcher.fits_role({"title": 20, "skill": 50}))
        self.assertFalse(matcher.fits_role({"title": 0, "skill": 80}))
        self.assertTrue(matcher.fits_role({"title": 0, "skill": 95}))

    def test_prune_removes_unrelated_unused_matches_only(self):
        tmp = tempfile.TemporaryDirectory()
        old = db.DB_PATH
        db.DB_PATH = Path(tmp.name) / "r.db"
        try:
            db.init_db()
            c = db.create_consultant({"name": "Raji", "email": "r@x.com"})
            db.upsert_resume(c["id"], "r.txt",
                             "RAJI\nTeamcenter PLM Administrator\nSkills: Teamcenter, SQL, Jira, Selenium, Java",
                             ["teamcenter", "sql", "jira", "selenium", "java"])
            good = db.insert_job({"source": "t", "source_id": "1", "title": "Teamcenter Administrator",
                                  "company": "A", "location": "Remote", "remote_flag": True,
                                  "description": "Teamcenter sql jira", "posted_at": ""})
            bad = db.insert_job({"source": "t", "source_id": "2", "title": "SDET Engineer",
                                 "company": "B", "location": "Remote", "remote_flag": True,
                                 "description": "selenium java sql jira testing playwright cypress appium", "posted_at": ""})
            db.insert_match(c["id"], good, 80.0, {"skill": 90, "title": 50, "location": 100, "recency": 0}, [])
            db.insert_match(c["id"], bad, 70.0, {"skill": 80, "title": 0, "location": 100, "recency": 0}, [])
            self.assertEqual(matcher.prune_misfits(), 1)
            ids = [m["job_id"] for m in db.list_matches()]
            self.assertEqual(ids, [good])
        finally:
            db.DB_PATH = old
            tmp.cleanup()


class AdzunaRemoteTests(unittest.TestCase):
    def setUp(self):
        from app.sources import adzuna
        self.adzuna = adzuna
        self.calls = []
        self.old_get = adzuna.get
        self.tmp = tempfile.TemporaryDirectory()
        self.old_db = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "z.db"
        db.init_db()

        outer = self

        class R:
            def __init__(self, results):
                self.results = results
            def raise_for_status(self):
                pass
            def json(self):
                return {"count": len(self.results), "results": self.results}

        def fake_get(url, params=None, **kw):
            outer.calls.append(dict(params))
            if "what_or" in params:
                return R(outer.first)
            return R(outer.second)
        adzuna.get = fake_get
        self.job = {"id": 1, "title": "QA Manager", "company": {"display_name": "X"},
                    "location": {"display_name": "Austin, TX"}, "description": "work remotely",
                    "redirect_url": "http://x", "created": "2026-10-01"}

    def tearDown(self):
        self.adzuna.get = self.old_get
        db.DB_PATH = self.old_db
        self.tmp.cleanup()

    def test_remote_search_leaves_where_out_and_asks_for_remote(self):
        self.first = [self.job]
        self.second = []
        jobs = self.adzuna.fetch({"title": "Senior QA Test Manager", "location": "Remote"},
                                 {"adzuna_app_id": "a", "adzuna_app_key": "b"})
        self.assertEqual(len(jobs), 1)
        self.assertNotIn("where", self.calls[0])
        self.assertIn("remote", self.calls[0]["what_or"])
        self.assertTrue(jobs[0]["remote_flag"])

    def test_remote_search_retries_without_remote_word_when_empty(self):
        self.first = []
        self.second = [self.job]
        jobs = self.adzuna.fetch({"title": "Teamcenter PLM", "location": "Remote"},
                                 {"adzuna_app_id": "a", "adzuna_app_key": "b"})
        self.assertEqual(len(jobs), 1)
        self.assertNotIn("what_or", self.calls[-1])
        self.assertEqual(db.usage_today("adzuna"), 2)

    def test_retry_drops_on_site_jobs(self):
        self.first = []
        onsite = dict(self.job, description="office based role", title="PLM Admin")
        self.second = [onsite]
        jobs = self.adzuna.fetch({"title": "Teamcenter PLM", "location": "Remote"},
                                 {"adzuna_app_id": "a", "adzuna_app_key": "b"})
        self.assertEqual(jobs, [])

    def test_remote_filter_error_falls_back_instead_of_failing(self):
        self.second = [self.job]
        orig = self.adzuna.get

        def flaky(url, params=None, **kw):
            if "what_or" in params:
                raise RuntimeError("503 Server Error")
            return orig(url, params=params, **kw)
        self.adzuna.get = flaky
        jobs = self.adzuna.fetch({"title": "QA", "location": "Remote"},
                                 {"adzuna_app_id": "a", "adzuna_app_key": "b"})
        self.assertEqual(len(jobs), 1)

    def test_few_results_trigger_a_shorter_title_search(self):
        titles = []
        job2 = dict(self.job, id=2, title="QA Test Manager")

        def fake(url, params=None, **kw):
            titles.append(params["what"])

            class R:
                def raise_for_status(self): pass
                def json(self_):
                    res = [self.job] if params["what"].startswith("Senior") else [job2]
                    return {"count": 1, "results": res}
            return R()
        self.adzuna.get = fake
        jobs = self.adzuna.fetch({"title": "Senior QA Test Manager", "location": "Remote"},
                                 {"adzuna_app_id": "a", "adzuna_app_key": "b"})
        self.assertIn("QA Test Manager", titles)
        self.assertEqual({j["source_id"] for j in jobs}, {"1", "2"})

    def test_loosen_title(self):
        self.assertEqual(self.adzuna.loosen_title("Senior QA Test Manager"), "QA Test Manager")
        self.assertEqual(self.adzuna.loosen_title("Teamcenter PLM"), "Teamcenter PLM")
        self.assertEqual(self.adzuna.loosen_title("Senior Engineer"), "")

    def test_error_text_hides_keys(self):
        from app.sources.common import safe_error
        msg = safe_error(Exception("503 for url: https://x/search/1?app_id=0b71&app_key=a458a9&what=QA"))
        self.assertNotIn("a458a9", msg)
        self.assertNotIn("0b71", msg)
        self.assertIn("what=QA", msg)

    def test_city_search_still_sends_the_city(self):
        self.first = [self.job]
        self.second = [self.job]
        self.adzuna.fetch({"title": "QA", "location": "Dallas, TX"},
                          {"adzuna_app_id": "a", "adzuna_app_key": "b"})
        self.assertEqual(self.calls[0]["where"], "Dallas, TX")
        self.assertNotIn("what_or", self.calls[0])



class CompanyTests(unittest.TestCase):
    """Three companies, each with its own login and its own database file:
    nobody can see another company's consultants, resumes or matches."""
    COMPANIES = {"X": {"username": "xuser", "password": "x-pass-1"},
                 "Y": {"username": "yuser", "password": "y-pass-2"},
                 "Z": {"username": "zuser", "password": "z-pass-3"}}

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._old = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "benchpilot.db"
        self._env = {k: os.environ.get(k) for k in
                     ("BENCHPILOT_COMPANIES", "BENCHPILOT_PASSWORD", "BENCHPILOT_USERNAME")}
        os.environ["BENCHPILOT_COMPANIES"] = json.dumps(self.COMPANIES)
        os.environ["BENCHPILOT_PASSWORD"] = "main-pass"
        os.environ["BENCHPILOT_USERNAME"] = "benchteam26"
        for p in db.all_paths():
            with db.use(p):
                db.init_db()
        self.paths = {c["name"]: c["path"] for c in db.companies()}

    def tearDown(self):
        for k, v in self._env.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        db.DB_PATH = self._old
        self.tmp.cleanup()

    def add(self, company, name, text):
        with db.use(self.paths[company] if company else None):
            c = db.create_consultant({"name": name, "target_roles": "SDET"})
            db.upsert_resume(c["id"], name + ".txt", text, skills.extract_skills(text)
                             if hasattr(skills, "extract_skills") else ["selenium"])
            return c

    def names(self, company):
        with db.use(self.paths[company] if company else None):
            return sorted(c["name"] for c in db.list_consultants())

    def test_secret_is_read(self):
        got = {c["name"]: c["username"] for c in db.companies()}
        self.assertEqual(got, {"X": "xuser", "Y": "yuser", "Z": "zuser"})
        self.assertEqual(len({str(c["path"]) for c in db.companies()}), 3)

    def test_bad_secret_means_no_companies(self):
        os.environ["BENCHPILOT_COMPANIES"] = "not json"
        self.assertEqual(db.companies(), [])

    def test_each_login_opens_only_its_own_file(self):
        self.assertEqual(db.login_target("xuser", "x-pass-1"), self.paths["X"])
        self.assertEqual(db.login_target("YUSER", "y-pass-2"), self.paths["Y"])
        self.assertEqual(db.login_target("zuser", "z-pass-3"), self.paths["Z"])
        self.assertIsNone(db.login_target("benchteam26", "main-pass"))
        # another company's password, the main password, or nothing: refused
        self.assertIs(db.login_target("xuser", "y-pass-2"), False)
        self.assertIs(db.login_target("xuser", "main-pass"), False)
        self.assertIs(db.login_target("benchteam26", "x-pass-1"), False)
        self.assertIs(db.login_target("someone", "x-pass-1"), False)
        self.assertIs(db.login_target("xuser", ""), False)

    def test_companies_do_not_see_each_other(self):
        self.add("X", "Asha X", "SDET with Selenium and Java")
        self.add("Y", "Bala Y", "QA engineer with Cypress")
        self.add(None, "Leha", "SDET with Playwright")
        self.assertEqual(self.names("X"), ["Asha X"])
        self.assertEqual(self.names("Y"), ["Bala Y"])
        self.assertEqual(self.names("Z"), [])
        self.assertEqual(self.names(None), ["Leha"])
        with db.use(self.paths["Z"]):
            self.assertEqual(db.consultants_with_resumes(), [])
            self.assertIsNone(db.get_resume(1))   # X's first consultant id: not here

    def test_jobs_are_shared_and_matched_inside_each_company(self):
        import collector
        self.add("X", "Asha X", "SDET automation Selenium Java TestNG API testing")
        job = {"source": "remotive", "source_id": "j1", "title": "SDET",
               "company": "Acme", "location": "Remote", "remote_flag": True,
               "url": "https://e.com/1", "description": "SDET Selenium Java TestNG API testing"}
        self.assertTrue(collector._save_job(job))
        for p in db.all_paths():
            with db.use(p):
                self.assertEqual(db.count_jobs(), 1)
        for p in db.all_paths():
            with db.use(p):
                matcher.run_all(threshold=1)
        with db.use(self.paths["X"]):
            self.assertTrue(db.list_matches())
        for other in ("Y", "Z"):
            with db.use(self.paths[other]):
                self.assertEqual(db.list_matches(), [])

    def test_searches_come_from_every_company(self):
        import collector
        self.add("X", "Asha X", "SDET automation Selenium")
        with db.use(self.paths["Y"]):
            c = db.create_consultant({"name": "Bala Y", "target_roles": "Data Engineer"})
            db.upsert_resume(c["id"], "b.txt", "Data Engineer Spark Python", ["spark"])
        queries, per_db = collector._prepare_all()
        titles = " ".join(q["title"] for q in queries).lower()
        self.assertIn("sdet", titles)
        self.assertIn("data engineer", titles)
        # each company's own part names only its own searches
        x = " ".join(q["title"] for q in per_db[str(self.paths["X"])]["queries"]).lower()
        self.assertNotIn("data engineer", x)
        self.assertEqual(per_db[str(self.paths["Z"])]["queries"], [])

    def test_api_counters_are_shared(self):
        with db.use(self.paths["X"]):
            db.usage_add("jsearch", 2)
        with db.use(self.paths["Y"]):
            self.assertEqual(db.usage_today("jsearch"), 2)
        self.assertEqual(db.usage_today("jsearch"), 2)

    def test_new_company_starts_with_existing_jobs(self):
        db.insert_job({"source": "remotive", "source_id": "old1", "title": "SDET",
                       "company": "Acme", "location": "Remote", "url": "u"})
        os.environ["BENCHPILOT_COMPANIES"] = json.dumps(
            dict(self.COMPANIES, W={"username": "wuser", "password": "w"}))
        w = [c for c in db.companies() if c["name"] == "W"][0]["path"]
        with db.use(w):
            db.init_db()
            self.assertEqual(db.count_jobs(), 1)
            self.assertEqual(db.list_consultants(), [])

    def test_each_company_is_backed_up_separately(self):
        from app import persist
        old = persist._last_sig
        persist._last_sig = {}
        try:
            store = FakeStore()
            self.add("X", "Asha X", "SDET")
            persist.backup_all(store)
            self.assertIn(persist.OBJECT, store.files)
            for name in ("x", "y", "z"):
                self.assertIn(f"benchpilot/company_{name}.db", store.files)
            self.paths["X"].unlink()                       # Republish wipes the disk
            self.assertTrue(persist.restore(store, self.paths["X"]))
            self.assertEqual(self.names("X"), ["Asha X"])
            self.assertEqual(self.names("Y"), [])
        finally:
            persist._last_sig = old

class StandingRulesTests(unittest.TestCase):
    """Per-consultant standing rules: parsing, storage, and rules_allow."""
    RESUME = ("Priya Nair\nQA Manager | Dallas, TX\n\nSUMMARY\nQA leader.\n\n"
              "TECHNICAL SKILLS\nLanguages: Java, SQL\nTools: Selenium, Jenkins\n")

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._old = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "t.db"
        db.init_db()

    def tearDown(self):
        db.DB_PATH = self._old
        self.tmp.cleanup()

    @staticmethod
    def _job(**kw):
        job = {"title": "QA Manager", "company": "Co", "location": "Dallas, TX",
               "remote_flag": 0, "description": "Requirements: Java, Selenium",
               "salary": ""}
        job.update(kw)
        return job

    def test_empty_rules_allow_everything(self):
        for consultant in ({}, {"rules": None}, {"rules": {}},
                           {"rules_json": ""}, {"rules_json": "not json"}):
            ok, _ = rules.rules_allow(consultant, self._job())
            self.assertTrue(ok, consultant)
        self.assertEqual(rules.parse_rules(None), rules.empty_rules())

    def test_title_allow_and_block(self):
        c = {"rules": {"titles": ["Data Engineer", "Java"]}}
        ok, _ = rules.rules_allow(c, self._job(title="Senior Java Developer"))
        self.assertTrue(ok)
        ok, reason = rules.rules_allow(c, self._job(title="QA Manager"))
        self.assertFalse(ok)
        self.assertIn("title", reason)

    def test_exclusion_keyword_wins_over_matching_title(self):
        c = {"rules": {"titles": ["QA"], "exclude_keywords": ["clearance"]}}
        job = self._job(title="QA Lead",
                        description="Active clearance required. Java, Selenium")
        ok, reason = rules.rules_allow(c, job)
        self.assertFalse(ok)
        self.assertIn("clearance", reason)

    def test_exclude_client_matches_company(self):
        c = {"rules": {"exclude_clients": ["acme"]}}
        ok, reason = rules.rules_allow(c, self._job(company="Acme Corp"))
        self.assertFalse(ok)
        self.assertIn("acme", reason.lower())
        ok, _ = rules.rules_allow(c, self._job(company="Other LLC"))
        self.assertTrue(ok)

    def test_remote_only(self):
        c = {"rules": {"remote_only": True}}
        ok, _ = rules.rules_allow(c, self._job(remote_flag=1, location=""))
        self.assertTrue(ok)
        ok, _ = rules.rules_allow(c, self._job(location="Remote - US"))
        self.assertTrue(ok)
        ok, reason = rules.rules_allow(c, self._job(location="Dallas, TX"))
        self.assertFalse(ok)
        self.assertIn("remote", reason)

    def test_locations_block_unlisted_but_never_remote(self):
        c = {"rules": {"locations": ["Houston"]}}
        ok, _ = rules.rules_allow(c, self._job(location="Houston, TX"))
        self.assertTrue(ok)
        ok, reason = rules.rules_allow(c, self._job(location="Chicago, IL"))
        self.assertFalse(ok)
        self.assertIn("location", reason)
        # a remote job is not blocked by a location list
        ok, _ = rules.rules_allow(c, self._job(location="Chicago, IL",
                                               remote_flag=1))
        self.assertTrue(ok)

    def test_min_rate(self):
        c = {"rules": {"min_rate": 120000}}
        ok, reason = rules.rules_allow(c, self._job(salary="$100,000 - $110,000"))
        self.assertFalse(ok)
        self.assertIn("min rate", reason)
        ok, _ = rules.rules_allow(c, self._job(salary="$130,000"))
        self.assertTrue(ok)
        ok, _ = rules.rules_allow(c, self._job(salary="$150k"))
        self.assertTrue(ok)
        # hourly figures are annualised before comparing
        ok, _ = rules.rules_allow(c, self._job(salary="$50/hr"))
        self.assertFalse(ok)
        ok, _ = rules.rules_allow(c, self._job(salary="$65 per hour"))
        self.assertTrue(ok)
        # missing or unreadable salary never blocks
        ok, _ = rules.rules_allow(c, self._job(salary=""))
        self.assertTrue(ok)
        ok, _ = rules.rules_allow(c, self._job(salary="competitive"))
        self.assertTrue(ok)

    def test_rules_round_trip_through_db(self):
        want = {"titles": ["Java Developer"], "locations": ["Houston, TX"],
                "remote_only": True, "min_rate": 120000,
                "exclude_keywords": ["clearance"], "exclude_clients": ["Acme"]}
        c = db.create_consultant({"name": "Priya Nair", "rules": want})
        self.assertEqual(c["rules"], want)
        self.assertEqual(db.get_consultant(c["id"])["rules"], want)
        self.assertEqual(db.list_consultants()[0]["rules"], want)
        # no rules sent -> the empty shape, allow-everything
        plain = db.create_consultant({"name": "No Rules"})
        self.assertEqual(plain["rules"], rules.empty_rules())
        # updating replaces the stored rules
        changed = dict(want, min_rate=90000, remote_only=False)
        got = db.update_consultant(c["id"], {"rules": changed})
        self.assertEqual(got["rules"], changed)
        # and a pre-existing row gains the parsed default too
        self.assertEqual(db.get_consultant(plain["id"])["rules"],
                         rules.empty_rules())


class AutopilotQueueTests(unittest.TestCase):
    """Tailor at/above the cutoff auto-files drafts into the apply queue
    (auto_queue_tailored), and standing rules gate every agent action."""
    RESUME = StandingRulesTests.RESUME

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self._old = db.DB_PATH
        db.DB_PATH = Path(self.tmp.name) / "t.db"
        db.init_db()

    def tearDown(self):
        matcher._jd_skills.cache_clear()
        db.DB_PATH = self._old
        self.tmp.cleanup()

    def _consultant(self, rules=None):
        data = {"name": "Priya Nair"}
        if rules is not None:
            data["rules"] = rules
        c = db.create_consultant(data)
        db.upsert_resume(c["id"], "r.txt", self.RESUME,
                         skills.extract_skills(self.RESUME))
        return c

    def _match(self, cid, title, score, company="Co", salary=""):
        job = {"source": "t", "source_id": title, "title": title,
               "company": company, "location": "", "remote_flag": 0,
               "url": "", "description": "Requirements: Java, Selenium, Kafka",
               "posted_at": "", "salary": salary, "employment_type": ""}
        jid = db.insert_job(job)
        return db.insert_match(cid, jid, score, {}, [])

    def test_tailor_queues_draft_with_resume_attached(self):
        c = self._consultant()
        self._match(c["id"], "QA Manager", 90.0)
        got = office._tailor(db.get_settings())
        self.assertEqual(got["drafted"], 1)
        self.assertEqual(got["queued"], 1)
        apps = db.list_applications()
        self.assertEqual(len(apps), 1)
        self.assertEqual(apps[0]["status"], "queued")
        self.assertEqual(apps[0]["tailored_resume_id"],
                         got["drafts"][0]["tailored_id"])
        self.assertFalse(got["queued_items"][0]["attached"])
        # a second run files nothing twice
        again = office._tailor(db.get_settings())
        self.assertEqual((again["drafted"], again["queued"]), (0, 0))
        self.assertEqual(len(db.list_applications()), 1)

    def test_auto_queue_off_keeps_drafts_only(self):
        db.set_setting("auto_queue_tailored", "0")
        c = self._consultant()
        self._match(c["id"], "QA Manager", 90.0)
        got = office._tailor(db.get_settings())
        self.assertEqual(got["drafted"], 1)
        self.assertEqual(got["queued"], 0)
        self.assertEqual(db.list_applications(), [])

    def test_backfill_attaches_existing_draft_to_queued_match(self):
        c = self._consultant()
        mid = self._match(c["id"], "QA Manager", 90.0)
        tid = db.insert_tailored(mid, "tailored text", "keyword-fallback", [])
        db.queue_application(mid)  # queued by hand, without the resume
        self.assertIsNone(db.list_applications()[0]["tailored_resume_id"])
        got = office._tailor(db.get_settings())
        self.assertEqual(got["drafted"], 0)      # the draft already existed
        self.assertEqual(got["queued"], 1)
        self.assertTrue(got["queued_items"][0]["attached"])
        self.assertEqual(db.list_applications()[0]["tailored_resume_id"], tid)

    def test_backfill_queues_unqueued_draft(self):
        c = self._consultant()
        mid = self._match(c["id"], "QA Manager", 90.0)
        db.insert_tailored(mid, "tailored text", "keyword-fallback", [])
        got = office._tailor(db.get_settings())
        self.assertEqual(got["queued"], 1)
        apps = db.list_applications()
        self.assertEqual(len(apps), 1)
        self.assertIsNotNone(apps[0]["tailored_resume_id"])

    def test_rule_blocked_match_is_neither_drafted_nor_queued(self):
        c = self._consultant(rules={"titles": ["Data Engineer"]})
        self._match(c["id"], "QA Manager", 95.0)
        got = office._tailor(db.get_settings())
        self.assertEqual(got["drafted"], 0)
        self.assertEqual(got["queued"], 0)
        self.assertEqual(got["skipped_by_rules"], 1)
        self.assertEqual(db.list_applications(), [])
        with db.get_conn() as conn:
            n = conn.execute("SELECT COUNT(*) AS n FROM tailored_resumes").fetchone()["n"]
        self.assertEqual(n, 0)

    def test_rule_blocked_backfill_is_not_attached(self):
        # the job description names Kafka; the consultant rules it out
        c = self._consultant(rules={"exclude_keywords": ["kafka"]})
        mid = self._match(c["id"], "QA Manager", 90.0)
        db.insert_tailored(mid, "tailored text", "keyword-fallback", [])
        db.queue_application(mid)
        got = office._tailor(db.get_settings())
        self.assertEqual(got["queued"], 0)
        self.assertGreaterEqual(got["skipped_by_rules"], 1)
        self.assertIsNone(db.list_applications()[0]["tailored_resume_id"])

    def test_rules_do_not_change_matching(self):
        # standing rules gate agent actions only: the match itself remains
        c = self._consultant(rules={"titles": ["Data Engineer"],
                                    "min_rate": 500000})
        mid = self._match(c["id"], "QA Manager", 95.0, salary="$10")
        self.assertEqual([m["id"] for m in db.list_matches()], [mid])


if __name__ == "__main__":
    unittest.main()
